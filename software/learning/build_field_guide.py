#!/usr/bin/env python3
"""Build the offline Word guide and figures using python-docx and markdown-it-py."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
FIGURES = ROOT / "artifacts/learning"
INK = "#20282d"
GREEN = "#087f71"
BLUE = "#276cc2"
RED = "#b64643"


def font(size: int, bold: bool = False):
    path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf"
    return ImageFont.truetype(path, size)


def canvas(title: str):
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 40), title, font=font(38, True), fill=INK)
    return image, draw


def figures() -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    image, draw = canvas("Three experiments and their boundaries")
    rows = [
        ("PHYSICAL", "Pi camera", "Jetson detector", "Controller + Pi", "Motor driver", "Historical floor trials; Mac transport dependency", GREEN),
        ("REAL2SIM", "Motor + gyro log", "Yaw integration", "Parameter fit", "Held out turns", "Measured yaw fit; not a reconstructed house", BLUE),
        ("SIMULATION", "Virtual sensors", "Policy or planner", "MuJoCo forces", "Next state", "Virtual sensing includes inputs the real rover lacks", RED),
    ]
    for row, (name, *rest) in enumerate(rows):
        a, b, c, d, note, colour = rest
        y = 140 + row * 240
        draw.text((65, y), name, font=font(26, True), fill=colour)
        for i, label in enumerate((a, b, c, d)):
            x = 65 + i * 385
            draw.text((x, y + 57), label, font=font(27), fill=INK)
            draw.line((x, y + 110, x + 295, y + 110), fill=colour, width=5)
            if i < 3:
                draw.text((x + 325, y + 53), ">", font=font(30), fill=INK)
        draw.text((65, y + 140), note, font=font(26), fill=INK)
    image.save(FIGURES / "system-boundaries.png")

    data = json.loads((ROOT / "artifacts/sim/2026-09-06-closed-loop-turn-validation.json").read_text())
    image, draw = canvas("Held out final turn angles")
    y0, y1, low, high = 160, 735, -210, 90
    to_y = lambda value: y1 - (value - low) / (high - low) * (y1 - y0)
    for value in range(-180, 91, 30):
        y = to_y(value)
        draw.line((160, y, 1520, y), fill="#dce1e3", width=2)
        draw.text((55, y - 15), str(value), font=font(25), fill=INK)
    draw.text((65, 105), "Degrees", font=font(24), fill=INK)
    keys = [("physical_degrees", "Physical", GREEN), ("legacy_sim_degrees", "Old sim", RED), ("fitted_sim_degrees", "Fitted sim", BLUE)]
    for i, turn in enumerate(data["turns"]):
        x = 320 + i * 460
        for offset, (key, _, colour) in enumerate(keys):
            bx = x - 95 + offset * 78
            value = turn[key]
            yy = to_y(value)
            draw.rectangle((bx, min(yy, to_y(0)), bx + 48, max(yy, to_y(0))), fill=colour)
            draw.text((bx - 8, yy - 33 if value >= 0 else yy + 5), f"{value:.1f}", font=font(22), fill=INK)
        draw.text((x - 42, 775), f"Turn {i + 1}", font=font(26), fill=INK)
    for i, (_, name, colour) in enumerate(keys):
        x = 290 + i * 385
        draw.rectangle((x, 844, x + 22, 866), fill=colour)
        draw.text((x + 35, 836), name, font=font(26), fill=INK)
    image.save(FIGURES / "turn-validation.png")

    data = json.loads((ROOT / "artifacts/sim/2026-09-08-controller-ablation.json").read_text())["summary"]
    image, draw = canvas("Safety filter ablation on 12 seeded worlds")
    names = ["Baseline", "Baseline + filter", "Network", "Network + filter"]
    for i, (name, values) in enumerate(zip(names, data.values())):
        y = 170 + i * 150
        draw.text((65, y), name, font=font(27, True), fill=INK)
        draw.rectangle((445, y, 445 + values["successes"] * 33, y + 30), fill=GREEN)
        draw.text((870, y - 5), f'{values["successes"]}/12 arrived', font=font(26), fill=INK)
        if values["contact_events"]:
            draw.rectangle((445, y + 51, 445 + values["contact_events"] * 12, y + 81), fill=RED)
        draw.text((1050, y + 43), f'{values["contact_events"]} contact events', font=font(26), fill=INK)
    draw.text((65, 820), "Separate scales: green = successes; red = contact events. Not physical trials.", font=font(25), fill=INK)
    image.save(FIGURES / "controller-comparison.png")


def hyperlink(paragraph, text: str, url: str):
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "245F8F")
    props.append(color)
    run.append(props)
    val = OxmlElement("w:t")
    val.text = text
    run.append(val)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, token):
    bold = italic = False
    url = None
    for child in token.children or []:
        if child.type == "strong_open": bold = True
        elif child.type == "strong_close": bold = False
        elif child.type == "em_open": italic = True
        elif child.type == "em_close": italic = False
        elif child.type == "link_open": url = child.attrGet("href")
        elif child.type == "link_close": url = None
        elif child.type in ("text", "code_inline", "softbreak", "hardbreak"):
            text = " " if child.type == "softbreak" else "\n" if child.type == "hardbreak" else child.content
            if url:
                hyperlink(paragraph, text, url)
            else:
                run = paragraph.add_run(text)
                run.bold, run.italic = bold, italic
                if child.type == "code_inline":
                    run.font.name = "DejaVu Sans Mono"
                    run.font.size = Pt(9)


def document(destination: Path) -> None:
    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse((ROOT / "docs/offline-field-guide.md").read_text())
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 27), ("Heading 1", 19), ("Heading 2", 14)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(12)
    for style in doc.styles:
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    doc.styles["Subtitle"].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    code = doc.styles.add_style("Source Code", 1)
    code.font.name, code.font.size = "DejaVu Sans Mono", Pt(8.5)
    code.paragraph_format.space_after = Pt(1)
    code.paragraph_format.line_spacing = 1.05
    code.paragraph_format.left_indent = Inches(0.08)
    footer = section.footer.paragraphs[0]
    footer.add_run("Frontier Rover | Offline field guide | ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    footer.style = "Caption"
    list_depth = 0
    ordered = False
    table = row = cell = None
    paragraph = None
    chapter = 0
    compact_chapter = False
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type == "heading_open":
            level = int(token.tag[1])
            text = tokens[i + 1].content
            if level == 2:
                compact_chapter = text == "Motor commands and the MPU"
            if level == 1:
                paragraph = doc.add_paragraph(text, "Title")
                doc.add_paragraph("Jack Coleman  |  September 2026", "Subtitle")
            else:
                paragraph = doc.add_paragraph(text, "Heading 1" if level == 2 else "Heading 2")
                if level == 2:
                    if chapter > 0:
                        paragraph.paragraph_format.page_break_before = True
                    chapter += 1
            i += 2
        elif token.type in ("bullet_list_open", "ordered_list_open"):
            list_depth += 1
            ordered = token.type == "ordered_list_open"
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            list_depth -= 1
        elif token.type == "paragraph_open":
            if cell is not None:
                paragraph = cell.paragraphs[0]
            else:
                paragraph = doc.add_paragraph(style=("List Number" if ordered else "List Bullet") if list_depth else "Normal")
                if compact_chapter:
                    paragraph.paragraph_format.space_after = Pt(3)
        elif token.type == "inline":
            images = [x for x in token.children or [] if x.type == "image"]
            if images:
                for item in images:
                    path = (ROOT / "docs" / item.attrGet("src")).resolve()
                    if path.is_file():
                        paragraph.add_run().add_picture(str(path), width=Inches(6.75))
                        paragraph.paragraph_format.keep_with_next = True
                        doc.add_paragraph(item.content, "Caption")
            else:
                if paragraph is None:
                    paragraph = doc.add_paragraph()
                inline(paragraph, token)
        elif token.type == "fence":
            lines = token.content.rstrip().splitlines()
            for n, line in enumerate(lines):
                p = doc.add_paragraph(line or " ", "Source Code")
                p.paragraph_format.keep_with_next = n < len(lines) - 1
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif token.type == "table_open":
            table = doc.add_table(rows=0, cols=0)
            table.style = "Table Grid"
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement(f"w:{edge}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
        elif token.type == "tr_open":
            if len(table.columns) == 0:
                count = 0
                for look in tokens[i + 1:]:
                    if look.type == "tr_close": break
                    if look.type in ("th_open", "td_open"): count += 1
                for _ in range(count): table.add_column(Inches(6.8 / count))
            row = table.add_row()
            cell = None
            col_index = 0
        elif token.type in ("th_open", "td_open"):
            cell = row.cells[col_index]
            col_index += 1
            cell.vertical_alignment = 1
            props = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "30383D" if token.type == "th_open" else
                        "F3F5F6" if len(table.rows) % 2 == 0 else "FFFFFF")
            props.append(shading)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "90")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(6)
            if token.type == "th_open":
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
        elif token.type in ("th_close", "td_close"):
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(255, 255, 255) if token.type == "th_close" else RGBColor(0, 0, 0)
                    if token.type == "th_close": run.bold = True
            cell = None
        elif token.type == "table_close":
            table = row = cell = None
            doc.add_paragraph()
        i += 1
    doc.core_properties.title = "Frontier Rover Machine Vision and Simulation Field Guide"
    doc.core_properties.author = "Jack Coleman"
    doc.core_properties.subject = "Project evidence, mechanisms, examples and offline exercises"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    print(destination)


if __name__ == "__main__":
    args = argparse.ArgumentParser()
    args.add_argument("--output", type=Path, required=True)
    parsed = args.parse_args()
    figures()
    document(parsed.output)
